#!/usr/bin/env python3
"""Export the workbook's player pool, league settings, and its own prices as JSON.

Usage:  python3 tools/export_players.py <path-to.xlsm> [outdir]

The engine in app/src/engine recomputes prices from the exported stat lines and
settings alone; the exported prices are the ground truth it is checked against.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

SKILL = ("QB", "RB", "WR", "TE")
# Stat columns are shared by RB/WR/TE; QB uses its own passing layout.
QB_STATS = {"passAtt": 7, "cmp": 8, "passYds": 9, "passTD": 10, "int": 11,
            "rushAtt": 12, "rushYds": 13, "rushTD": 14, "fum": 15}
SKILL_STATS = {"rushAtt": 9, "rushYds": 10, "rushTD": 11,
               "rec": 12, "recYds": 13, "recTD": 14, "fum": 15}
ZERO = {"passAtt": 0, "cmp": 0, "passYds": 0, "passTD": 0, "int": 0, "rushAtt": 0,
        "rushYds": 0, "rushTD": 0, "rec": 0, "recYds": 0, "recTD": 0, "fum": 0}


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def num(v):
    return round(v, 4) if isinstance(v, (int, float)) else 0


def scoring(li):
    """The four per-position scoring tables, keyed the way the engine wants them."""
    def block(col, rows):
        return {k: num(li.cell(row=r, column=col).value) for k, r in rows.items()}
    qb = block(10, {"passTD": 2, "passYds": 3, "int": 4, "rushAtt": 5, "rushTD": 6,
                    "rushYds": 7, "fum": 8, "passAtt": 9, "cmp": 10})
    rb = block(10, {"rushAtt": 12, "rushTD": 13, "rushYds": 14, "rec": 15,
                    "recTD": 16, "recYds": 17, "fum": 18})
    wr = block(14, {"rushAtt": 2, "rushTD": 3, "rushYds": 4, "rec": 5,
                    "recTD": 6, "recYds": 7, "fum": 8})
    te = block(14, {"rushAtt": 12, "rushTD": 13, "rushYds": 14, "rec": 15,
                    "recTD": 16, "recYds": 17, "fum": 18})
    return {"QB": ZERO | qb, "RB": ZERO | rb, "WR": ZERO | wr, "TE": ZERO | te}


def league(li):
    return {
        "teams": li["F2"].value,
        "budget": li["F3"].value,
        "rosterSize": li["F4"].value,
        "benchSize": li["F5"].value,
        "starters": {"QB": li["D3"].value, "RB": li["D4"].value, "WR": li["D5"].value,
                     "TE": li["D6"].value, "FLEX": li["D7"].value,
                     "FLEX2": li["D8"].value, "K": li["D9"].value, "DEF": li["D10"].value},
        "flexType": li["F8"].value,
        "flexType2": li["F9"].value,
        "starterPct": li["D15"].value,
        "method": li["H7"].value,
        "allowNegative": li["H6"].value == "Yes",
        "marketSite": li["F7"].value,
        "scoring": scoring(li),
    }


def market(wb):
    ws = wb["DefaultAuctionValues"]
    out = {}
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        out[slug(name)] = {"yahoo": num(ws.cell(r, 2).value),
                           "espn": num(ws.cell(r, 3).value),
                           "nffc": num(ws.cell(r, 4).value)}
    return out


def main():
    src = Path(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "app/data")
    out.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, data_only=True)

    players, expected = [], {}
    prices = market(wb)
    for pos in SKILL:
        ws = wb[pos]
        cols = QB_STATS if pos == "QB" else SKILL_STATS
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 2).value
            if not name:
                continue
            pid = slug(name)
            players.append({
                "id": pid, "name": name, "pos": pos,
                "team": ws.cell(r, 6).value, "bye": ws.cell(r, 5).value,
                "stats": ZERO | {k: num(ws.cell(r, c).value) for k, c in cols.items()},
                "market": prices.get(pid, {"yahoo": 0, "espn": 0, "nffc": 0}),
            })
            expected[pid] = round(ws.cell(r, 20).value, 6)
    for pos in ("K", "DEF"):
        ws = wb[pos]
        for r in range(2, ws.max_row + 1):
            name = ws.cell(r, 2).value
            if not name:
                continue
            pid = slug(name)
            players.append({
                "id": pid, "name": name, "pos": pos,
                "team": ws.cell(r, 6).value, "bye": ws.cell(r, 5).value,
                "fpts": num(ws.cell(r, 10).value), "stats": None,
                "market": prices.get(pid, {"yahoo": 0, "espn": 0, "nffc": 0}),
            })
            expected[pid] = round(ws.cell(r, 20).value, 6)

    (out / "players.json").write_text(json.dumps(players, indent=1))
    (out / "league.json").write_text(json.dumps(league(wb["LeagueInfo"]), indent=1))
    (out / "expected_prices.json").write_text(json.dumps(expected, indent=1))
    print(f"{len(players)} players -> {out}/players.json")


if __name__ == "__main__":
    main()
